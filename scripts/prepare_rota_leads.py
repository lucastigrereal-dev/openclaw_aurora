from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl

MISSING_MARKERS = {
    "", "nao encontrado", "não encontrado", "nao localizado publicamente",
    "não localizado publicamente", "nao confirmado", "não confirmado", "none",
}
TEAM_ID = "9e7c50c1-1cbf-481b-ae72-c309604ec06c"
SOURCE_NAME = "Rota_das_Aguas_Atlas_SUPER_MASTER"
CANONICAL_NAMES = {
    "bar do tio", "bistro d franca", "cafe das fontes",
    "cantina mexicana mi hermana", "emporio das aguas", "rei do peixe",
    "restaurante do lago", "sorveteria rocha", "villa tardivelli",
    "zuleika s doces", "grande hotel sao pedro", "hotel aguas de sao pedro",
    "hotel avenida charme", "hotel jerubiacaba", "hotel portal das aguas",
    "hotel recanto das aguas", "ls villas hotel spa",
    "pousada caminho das aguas", "pousada por do sol",
    "pousada vale das aguas", "cafe emporio parque das aguas",
}


def plain(value: Any) -> str:
    return str(value or "").strip()


def folded(value: Any) -> str:
    text = unicodedata.normalize("NFD", plain(value))
    return "".join(ch for ch in text if unicodedata.category(ch) != "Mn").lower()


def name_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", folded(value)).strip()


def clean(value: Any) -> str:
    text = plain(value)
    return "" if folded(text) in MISSING_MARKERS else text


def digits(value: Any) -> str:
    return re.sub(r"\D+", "", clean(value))


def handle(value: Any) -> str:
    return clean(value).lstrip("@").lower()


def safe_id(value: str) -> str:
    normalized = folded(value)
    return re.sub(r"[^a-z0-9]+", "-", normalized).strip("-")[:64]


def temperature(score: Any, potential: Any) -> str:
    try:
        numeric = float(score or 0)
    except (TypeError, ValueError):
        numeric = 0
    if numeric >= 9 or "alto" in folded(potential):
        return "quente"
    if numeric >= 6:
        return "morno"
    return "frio"


def load_rows(source: Path) -> list[dict[str, Any]]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    values = list(sheet.iter_rows(values_only=True))
    headers = [plain(value) for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:] if any(row)]


def choose_identity(row: dict[str, Any]) -> tuple[str, str]:
    phone = digits(row.get("WhatsApp")) or digits(row.get("Telefone 1"))
    instagram = handle(row.get("Instagram"))
    if phone:
        return "phone", phone
    if instagram:
        return "instagram", instagram
    return "", ""


def make_record(row: dict[str, Any], now: str) -> dict[str, Any]:
    identity_type, identity = choose_identity(row)
    name = clean(row.get("Nome"))
    phone = digits(row.get("WhatsApp")) or digits(row.get("Telefone 1"))
    instagram = handle(row.get("Instagram"))
    verification = "phone_confirmed_from_master" if phone else "instagram_accessible_needs_validation"
    data = {
        "name": name,
        "email": "",
        "phone": phone,
        "handle": instagram,
        "contactVerification": verification,
        "origin": SOURCE_NAME,
        "interest": clean(row.get("Categoria")),
        "temperature": temperature(row.get("Score"), row.get("Potencial")),
        "status": "novo",
        "value": 0,
        "owner": "Cláudia",
        "nextAction": "Primeiro contato da Cláudia" if phone else "Validar perfil e enviar DM",
        "notes": clean(row.get("Resumo")),
        "city": clean(row.get("Cidade")),
        "businessUnitKey": "rota_das_aguas_sp",
        "importedAt": now,
        "updatedAt": now,
    }
    record_id = f"{TEAM_ID}:leads:lead-{safe_id(identity or name)}"
    return {"id": record_id, "team_id": TEAM_ID, "kind": "leads", "data": data}


def write_csv(path: Path, rows: list[dict[str, Any]], headers: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def prepare(source: Path, output_dir: Path) -> dict[str, Any]:
    now = datetime.now(timezone.utc).isoformat()
    rows = load_rows(source)
    rows.append({
        "Nome": "Café & Empório Parque das Águas",
        "Cidade": "Águas de São Pedro",
        "Categoria": "Negócio Leve",
        "Resumo": "Ambiente premium",
        "Telefone 1": "(19) 99628-5505",
        "WhatsApp": "(19) 99628-5505",
        "Instagram": "",
        "Potencial": "Alto",
        "Score": 7,
    })
    city_rows = [
        row for row in rows
        if folded(row.get("Cidade")) == "aguas de sao pedro"
        and name_key(row.get("Nome")) in CANONICAL_NAMES
    ]
    ready: list[dict[str, Any]] = []
    pending: list[dict[str, Any]] = []
    conflicts: list[dict[str, Any]] = []
    by_identity: dict[str, list[dict[str, Any]]] = {}

    for row in city_rows:
        identity_type, identity = choose_identity(row)
        if not identity:
            pending.append(row)
            continue
        key = f"{identity_type}:{identity}"
        by_identity.setdefault(key, []).append(row)

    for key, grouped in by_identity.items():
        distinct_names = {name_key(row.get("Nome")) for row in grouped}
        if len(distinct_names) > 1:
            for row in grouped:
                conflicts.append({**row, "Chave_Conflito": key})
            continue
        ready.append(grouped[0])

    output_dir.mkdir(parents=True, exist_ok=True)
    records = [make_record(row, now) for row in ready]
    import_rows = []
    for record in records:
        data = record["data"]
        import_rows.append({
            "Nome": data["name"], "Telefone": data["phone"],
            "Instagram": data["handle"], "Origem do lead": data["origin"],
            "Interesse": data["interest"], "Temperatura": data["temperature"],
            "Status": data["status"], "Responsavel": data["owner"],
            "Proxima acao": data["nextAction"], "Observacao": data["notes"],
        })

    headers = list(import_rows[0].keys()) if import_rows else ["Nome"]
    write_csv(output_dir / "aguas_de_sao_pedro_lote_01.csv", import_rows, headers)
    write_csv(output_dir / "pendentes_pesquisa.csv", pending, list(pending[0].keys()) if pending else ["Nome"])
    write_csv(output_dir / "conflitos_identidade.csv", conflicts, list(conflicts[0].keys()) if conflicts else ["Nome"])
    (output_dir / "crm_records.json").write_text(
        json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    report = {
        "source": str(source),
        "generated_at": now,
        "source_rows": len(rows),
        "city_rows": len(city_rows),
        "ready": len(ready),
        "pending_research": len(pending),
        "identity_conflicts": len(conflicts),
        "confirmed_phone_or_whatsapp": sum(1 for row in ready if digits(row.get("WhatsApp")) or digits(row.get("Telefone 1"))),
        "instagram_only": sum(1 for row in ready if not (digits(row.get("WhatsApp")) or digits(row.get("Telefone 1"))) and handle(row.get("Instagram"))),
        "names_ready": [clean(row.get("Nome")) for row in ready],
        "names_pending": [clean(row.get("Nome")) for row in pending],
    }
    (output_dir / "quality_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepara lote de leads da Rota das Águas")
    parser.add_argument(
        "--source",
        type=Path,
        default=Path(r"C:\Users\lucas\Downloads\Rota_das_Aguas_Atlas_AUDITORIA_FINAL_ENRIQUECIDA.xlsx"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(r"C:\Users\lucas\Downloads\rota-leads-go-live"),
    )
    args = parser.parse_args()
    if not args.source.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {args.source}")
    report = prepare(args.source, args.output)
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
